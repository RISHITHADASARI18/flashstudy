from dataclasses import dataclass
from pathlib import Path
import fitz
from docx import Document
from pptx import Presentation

try:
    import pytesseract
    from PIL import Image
except ImportError:
    pytesseract = None
    Image = None

@dataclass
class ExtractedUnit:
    position: int
    source_type: str
    source_number: int | None
    source_label: str
    text: str

SUPPORTED_TYPES = {
    ".pdf":"pdf",".doc":"doc",".docx":"docx",".ppt":"ppt",".pptx":"pptx",
    ".txt":"text",".md":"text",".jpg":"image",".jpeg":"image",".png":"image",".webp":"image",
}
SUPPORTED_MIME_TYPES = {
    "application/pdf":"pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document":"docx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation":"pptx",
    "text/plain":"text","text/markdown":"text",
    "image/jpeg":"image","image/png":"image","image/webp":"image",
}

def detect_material_type(filename: str, mime_type: str) -> str | None:
    extension_type = SUPPORTED_TYPES.get(Path(filename).suffix.lower())
    mime_type_value = SUPPORTED_MIME_TYPES.get(mime_type.lower())
    if extension_type and mime_type_value and extension_type != mime_type_value:
        return None
    return extension_type or mime_type_value

def extract_pdf(path: Path) -> list[ExtractedUnit]:
    units=[]
    with fitz.open(path) as doc:
        for i,page in enumerate(doc):
            text=page.get_text("text").strip()
            if text: units.append(ExtractedUnit(i,"page",i+1,f"Page {i+1}",text))
    return units

def extract_docx(path: Path) -> list[ExtractedUnit]:
    units=[]
    doc=Document(path)
    for i,p in enumerate(doc.paragraphs):
        text=p.text.strip()
        if text: units.append(ExtractedUnit(i,"paragraph",i+1,f"Paragraph {i+1}",text))
    return units

def extract_pptx(path: Path) -> list[ExtractedUnit]:
    units=[]
    deck=Presentation(path)
    for i,slide in enumerate(deck.slides):
        parts=[shape.text.strip() for shape in slide.shapes if hasattr(shape,"text") and shape.text.strip()]
        text="\n".join(parts).strip()
        if text: units.append(ExtractedUnit(i,"slide",i+1,f"Slide {i+1}",text))
    return units

def extract_text(path: Path) -> list[ExtractedUnit]:
    text=path.read_text(encoding="utf-8",errors="replace").strip()
    return [ExtractedUnit(0,"document",None,"Document",text)] if text else []

def extract_image(path: Path) -> list[ExtractedUnit]:
    if pytesseract is None or Image is None:
        raise RuntimeError("Image OCR dependencies are not available.")
    text=pytesseract.image_to_string(Image.open(path)).strip()
    return [ExtractedUnit(0,"image",None,"Image OCR",text)] if text else []

def extract_content(path: Path, material_type: str) -> list[ExtractedUnit]:
    if material_type=="pdf": return extract_pdf(path)
    if material_type=="docx": return extract_docx(path)
    if material_type=="pptx": return extract_pptx(path)
    if material_type=="text": return extract_text(path)
    if material_type=="image": return extract_image(path)
    raise RuntimeError(f"Extraction for {material_type.upper()} is not available yet.")
